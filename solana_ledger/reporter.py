"""
Excel report generator.

Sheet layout
────────────
  Summary          – wallet overview table + all external (non-intra) txns
                     + vote-fee totals by wallet/month
  <Wallet label>   – one sheet per wallet
                     • non-vote transactions, individual rows with running balance
                     • vote transactions aggregated by month (identity accounts
                       can have millions, so individual rows are impractical)
"""

import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .processor import LedgerEntry

# ── style constants ───────────────────────────────────────────────────────────

_BLUE_FILL = PatternFill("solid", fgColor="1F497D")
_LIGHT_BLUE = PatternFill("solid", fgColor="DCE6F1")
_GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
_WHITE_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_SOL_FMT = "#,##0.000000000"
_DATE_FMT = "YYYY-MM-DD HH:MM:SS"


# ── internal helpers ──────────────────────────────────────────────────────────

def _header_row(ws, row: int, headers: List[str]):
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = _WHITE_FONT
        c.fill = _BLUE_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(ws, max_width: int = 50):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        best = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(best + 3, max_width)


def _safe_name(label: str) -> str:
    name = re.sub(r'[\\/*?:\[\]]', "-", label)
    return name[:31]


def _sol(value: float) -> float | None:
    return round(value, 9) if value else None


# ── wallet sheet ──────────────────────────────────────────────────────────────

def _write_wallet_sheet(ws, label: str, entries: List[LedgerEntry]):
    ws.freeze_panes = "A2"

    non_votes = [e for e in entries if not e.is_vote]
    votes = [e for e in entries if e.is_vote]

    headers = [
        "Date (UTC)", "Signature", "Type", "Description",
        "Counterparty", "SOL In", "SOL Out", "Fee Paid (SOL)",
        "Running Balance (SOL)", "Intra-Wallet?",
    ]
    _header_row(ws, 1, headers)

    row = 2
    running = 0.0

    for e in non_votes:
        running += e.sol_in - e.sol_out
        fill = _LIGHT_BLUE if e.is_intra else None
        vals = [
            e.date.replace(tzinfo=None),   # Excel doesn't store tz
            e.signature,
            e.tx_type,
            e.description,
            e.counterparty,
            _sol(e.sol_in),
            _sol(e.sol_out),
            _sol(e.fee_sol),
            round(running, 9),
            "Yes" if e.is_intra else "",
        ]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            if fill:
                c.fill = fill
            if col == 1:
                c.number_format = _DATE_FMT
            if col in (6, 7, 8, 9):
                c.number_format = _SOL_FMT
        row += 1

    # ── vote aggregation section ──────────────────────────────────────────
    if votes:
        row += 1
        title_cell = ws.cell(row=row, column=1, value="VOTE TRANSACTIONS – AGGREGATED BY MONTH")
        title_cell.font = Font(bold=True, size=11)
        title_cell.fill = _GREY_FILL
        row += 1

        vote_headers = ["Month", "Vote Count", "Total Fees Paid (SOL)", "Avg Fee (SOL)"]
        for col, h in enumerate(vote_headers, 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = _BOLD
            c.fill = _GREY_FILL
        row += 1

        monthly: Dict[str, Dict] = defaultdict(lambda: {"count": 0, "fees": 0.0})
        for e in votes:
            key = e.date.strftime("%Y-%m")
            monthly[key]["count"] += 1
            monthly[key]["fees"] += e.sol_out

        for month in sorted(monthly):
            m = monthly[month]
            avg = m["fees"] / m["count"] if m["count"] else 0.0
            ws.cell(row=row, column=1, value=month)
            ws.cell(row=row, column=2, value=m["count"])
            c = ws.cell(row=row, column=3, value=round(m["fees"], 9))
            c.number_format = _SOL_FMT
            c = ws.cell(row=row, column=4, value=round(avg, 9))
            c.number_format = _SOL_FMT
            row += 1

        # Totals row
        ws.cell(row=row, column=1, value="TOTAL").font = _BOLD
        ws.cell(row=row, column=2, value=len(votes)).font = _BOLD
        total_c = ws.cell(row=row, column=3, value=round(sum(e.sol_out for e in votes), 9))
        total_c.font = _BOLD
        total_c.number_format = _SOL_FMT

    _auto_width(ws)


# ── summary sheet ─────────────────────────────────────────────────────────────

def _write_summary_sheet(
    ws,
    entries_by_wallet: Dict[str, List[LedgerEntry]],
    generated_at: datetime,
):
    row = 1
    ws.cell(row=row, column=1, value="Solana Validator Ledger").font = Font(bold=True, size=14)
    row += 1
    ws.cell(
        row=row, column=1,
        value=f"Generated: {generated_at.strftime('%Y-%m-%d %H:%M:%S UTC')}",
    ).font = Font(italic=True)
    row += 2

    # ── Wallet overview ───────────────────────────────────────────────────
    ws.cell(row=row, column=1, value="WALLET OVERVIEW").font = Font(bold=True, size=12)
    row += 1

    overview_hdrs = [
        "Label", "Address", "Type",
        "Total SOL In", "Total SOL Out", "Net SOL",
        "Vote Fees (SOL)", "Tx Count",
    ]
    _header_row(ws, row, overview_hdrs)
    row += 1

    for address, entries in entries_by_wallet.items():
        if not entries:
            continue
        e0 = entries[0]
        total_in = sum(e.sol_in for e in entries)
        total_out = sum(e.sol_out for e in entries)
        vote_fees = sum(e.sol_out for e in entries if e.is_vote)

        vals = [
            e0.wallet_label, address,
            next((w for w in [e0.wallet_address] if True), ""),
        ]
        ws.cell(row=row, column=1, value=e0.wallet_label)
        ws.cell(row=row, column=2, value=address)

        # find type from config (we don't have it here, so leave blank)
        ws.cell(row=row, column=3, value="")

        for col, val in [(4, total_in), (5, total_out), (6, total_in - total_out), (7, vote_fees)]:
            c = ws.cell(row=row, column=col, value=round(val, 9))
            c.number_format = _SOL_FMT
        ws.cell(row=row, column=8, value=len(entries))
        row += 1

    row += 2

    # ── External transactions (no intra-wallet, no votes) ─────────────────
    ws.cell(
        row=row, column=1,
        value="EXTERNAL TRANSACTIONS  (intra-wallet transfers excluded)",
    ).font = Font(bold=True, size=12)
    row += 1

    ext_hdrs = [
        "Date (UTC)", "Signature", "Wallet", "Type",
        "Counterparty", "SOL In", "SOL Out", "Fee (SOL)",
    ]
    _header_row(ws, row, ext_hdrs)
    row += 1

    seen: set = set()
    for entries in entries_by_wallet.values():
        for e in entries:
            if e.is_intra or e.is_vote:
                continue
            if e.signature in seen:
                continue
            seen.add(e.signature)

            ws.cell(row=row, column=1, value=e.date.replace(tzinfo=None)).number_format = _DATE_FMT
            ws.cell(row=row, column=2, value=e.signature)
            ws.cell(row=row, column=3, value=e.wallet_label)
            ws.cell(row=row, column=4, value=e.tx_type)
            ws.cell(row=row, column=5, value=e.counterparty)
            for col, val in [(6, e.sol_in), (7, e.sol_out), (8, e.fee_sol)]:
                c = ws.cell(row=row, column=col, value=_sol(val))
                c.number_format = _SOL_FMT
            row += 1

    row += 2

    # ── Vote fees by wallet and month ─────────────────────────────────────
    ws.cell(row=row, column=1, value="VOTE FEES SUMMARY BY WALLET AND MONTH").font = Font(bold=True, size=12)
    row += 1

    vf_hdrs = ["Wallet", "Month", "Vote Count", "Total Fees (SOL)"]
    _header_row(ws, row, vf_hdrs)
    row += 1

    for address, entries in entries_by_wallet.items():
        votes = [e for e in entries if e.is_vote]
        if not votes:
            continue
        label = votes[0].wallet_label
        monthly: Dict[str, Dict] = defaultdict(lambda: {"count": 0, "fees": 0.0})
        for e in votes:
            k = e.date.strftime("%Y-%m")
            monthly[k]["count"] += 1
            monthly[k]["fees"] += e.sol_out

        for month in sorted(monthly):
            m = monthly[month]
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=month)
            ws.cell(row=row, column=3, value=m["count"])
            c = ws.cell(row=row, column=4, value=round(m["fees"], 9))
            c.number_format = _SOL_FMT
            row += 1

    _auto_width(ws)


# ── public entry point ────────────────────────────────────────────────────────

def generate_report(
    entries_by_wallet: Dict[str, List[LedgerEntry]],
    wallet_labels: Dict[str, str],
    wallet_types: Dict[str, str],
    output_path: str,
) -> str:
    generated_at = datetime.now(tz=timezone.utc)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Summary first
    ws_sum = wb.create_sheet("Summary", 0)
    _write_summary_sheet(ws_sum, entries_by_wallet, generated_at)

    # Per-wallet sheets
    for address, entries in entries_by_wallet.items():
        label = wallet_labels.get(address, address[:8])
        ws = wb.create_sheet(_safe_name(label))
        _write_wallet_sheet(ws, label, entries)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
